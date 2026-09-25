"""Independent LOOCV screen of 3–7RM bench records vs body dimensions.

Source: https://zenodo.org/records/3253731 (CC BY 4.0). Download the two
small workbooks into .cache/research before running; do not redistribute rows.
This experiment cannot validate single-muscle size or force, which the source
did not measure.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from urllib.request import urlopen

import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
CACHE = ROOT / ".cache" / "research"
FILES = {
    "Data.xlsx": "88e52fc382e53c033a8771e8eb2c1df2",
    "ScalingData_all_measure_variables.xlsx": "e44b34e6593a5af512ecb0d0a2bc7352",
}
TARGETS = ("CHEST_BRTH", "CHEST_DEPTH", "BIDELTOID_BRTH", "BIACROMIAL_BRTH", "WRIST_CIRC-STYLION")


def source_file(name: str) -> Path:
    CACHE.mkdir(parents=True, exist_ok=True)
    path = CACHE / name
    if not path.exists():
        url = f"https://zenodo.org/api/records/3253731/files/{name}/content"
        path.write_bytes(urlopen(url, timeout=30).read())
    actual = hashlib.md5(path.read_bytes()).hexdigest()
    if actual != FILES[name]:
        raise ValueError(f"Unexpected MD5 for {name}: {actual}")
    return path


def fit_predict(train_x: np.ndarray, train_y: np.ndarray, test_x: np.ndarray) -> float:
    """Fixed ridge strength; all centering and scaling is train-fold only."""
    center = train_x.mean(axis=0)
    scale = train_x.std(axis=0)
    scale[scale < 1e-10] = 1
    x = (train_x - center) / scale
    z = (test_x - center) / scale
    y_center = train_y.mean()
    alpha = 10.0
    beta = np.linalg.solve(x.T @ x + alpha * np.eye(x.shape[1]), x.T @ (train_y - y_center))
    return float(y_center + z @ beta)


def main() -> None:
    strength = list(load_workbook(source_file("Data.xlsx"), read_only=True, data_only=True)["3-7RM"].values)
    size_rows = list(load_workbook(source_file("ScalingData_all_measure_variables.xlsx"), read_only=True, data_only=True)["Scal_dat"].values)
    size_header = {name: i for i, name in enumerate(size_rows[0])}
    sizes = {row[0]: row for row in size_rows[1:] if isinstance(row[0], str)}
    records = []
    for row in strength[2:]:
        sid = row[0]
        if not isinstance(sid, str) or sid not in sizes:
            continue
        s = sizes[sid]
        required = [row[1], row[2], row[3], row[4], s[1], s[2]] + [s[size_header[t]] for t in TARGETS]
        if any(not isinstance(v, (int, float)) for v in required):
            continue
        records.append((row, s))
    if len(records) != 28:
        raise ValueError(f"Expected 28 matched complete subjects; got {len(records)}")
    body = np.array([[r[0][1] * 100, r[0][2]] for r in records], dtype=float)
    body_demographic = np.array([[r[0][1] * 100, r[0][2], r[1][1] == 1, r[1][2]] for r in records], dtype=float)
    bench = np.array([[r[0][3], r[0][4]] for r in records], dtype=float)
    output = {"source": "https://zenodo.org/records/3253731", "n": len(records), "method": "LOOCV; ridge alpha 10; fold-only standardization", "targets": {}}
    for target in TARGETS:
        y = np.array([s[size_header[target]] * 100 for _, s in records], dtype=float)
        output["targets"][target] = {"mean_cm": round(float(y.mean()), 2), "std_cm": round(float(y.std()), 2)}
        for label, base in (("height_mass", body), ("height_mass_sex_age", body_demographic)):
            errors = []
            with_bench_errors = []
            for i in range(len(y)):
                mask = np.arange(len(y)) != i
                base_pred = fit_predict(base[mask], y[mask], base[i])
                extended = np.column_stack((base, bench))
                bench_pred = fit_predict(extended[mask], y[mask], extended[i])
                errors.append(abs(base_pred - y[i]))
                with_bench_errors.append(abs(bench_pred - y[i]))
            individual_delta = np.array(with_bench_errors) - np.array(errors)
            rng = np.random.default_rng(3253731)
            resampled = rng.choice(individual_delta, size=(20000, len(y)), replace=True).mean(axis=1)
            interval = np.quantile(resampled, [0.025, 0.975])
            output["targets"][target][label] = {
                "baseline_mae_cm": round(float(np.mean(errors)), 3),
                "with_bench_mae_cm": round(float(np.mean(with_bench_errors)), 3),
                "delta_mae_cm": round(float(np.mean(with_bench_errors) - np.mean(errors)), 3),
                "delta_95pct_bootstrap_ci_cm": [round(float(v), 3) for v in interval],
                "people_improved": int(sum(a > b for a, b in zip(errors, with_bench_errors))),
            }
    print(json.dumps(output, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
