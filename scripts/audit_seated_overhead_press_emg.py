"""Audit eight-person seated overhead-press phase EMG source workbook.

Source: Coratella et al. (2022), Zenodo 5744051, CC BY 4.0. The workbook
contains per-person phase-averaged normalized EMG, not raw time-series EMG or
standing barbell press kinematics/forces. Output is research-only.
"""

import argparse
import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]

SOURCE_MD5 = "e8b6c718179c993b0817208341e84130"
SOURCE = ROOT / ".cache/overhead-press-emg/Raw data_new.xlsx"
OUTPUT = ROOT / "docs/data/seated_overhead_press_emg_audit.json"
SHEETS = {
    "pectoralis_clavicular": "GranPettoraleClaveare%",
    "triceps_lateral": "TricipiteCapoLaterale%",
    "trapezius_upper": "TrapezioSuperiore%",
    "deltoid_anterior": "DeltoideAnteriore%",
    "deltoid_middle": "DeltoideMedio%",
    "deltoid_posterior": "DeltoidePosteriore%",
}
EXPECTED_HEADINGS = (
    "Lento_Avanti_CONC", "Lento_Avanti_ECC", "Lento_Dietro_CONC",
    "Lento_Dietro_ECC", "Lento_Avanti_MultiPower_CONC",
    "Lento_Avanti_MultiPower_ECC",
)


def rounded(value):
    return round(float(value), 2)


def paired(a, b):
    difference = a - b
    return {
        "medianDifferencePercentagePoints": rounded(np.median(difference)),
        "meanDifferencePercentagePoints": rounded(np.mean(difference)),
        "peopleHigher": int(np.sum(difference > 0)),
        "peopleLower": int(np.sum(difference < 0)),
        "peopleTied": int(np.sum(difference == 0)),
    }


def inspect(path):
    actual_md5 = hashlib.md5(path.read_bytes()).hexdigest()
    if actual_md5 != SOURCE_MD5:
        raise ValueError(f"Source workbook MD5 mismatch: {actual_md5}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    if set(workbook.sheetnames) != set(SHEETS.values()):
        raise ValueError(f"Unexpected sheets: {workbook.sheetnames}")
    muscles = {}
    for muscle, sheet_name in SHEETS.items():
        rows = list(workbook[sheet_name].values)
        if tuple(rows[0][:6]) != EXPECTED_HEADINGS:
            raise ValueError(f"Unexpected phase/condition headings: {sheet_name}")
        # Participant records are rows 2-9. Rows 12-13 are the source's
        # calculated mean and sample SD; other trailing cells are annotations.
        data = np.asarray([row[:8] for row in rows[1:9]], dtype=float)
        if data.shape != (8, 8) or not np.all(np.isfinite(data)):
            raise ValueError(f"Expected eight complete paired records: {sheet_name}")
        source_mean = np.asarray(rows[11][:8], dtype=float)
        source_sd = np.asarray(rows[12][:8], dtype=float)
        if (not np.allclose(data.mean(axis=0), source_mean, atol=1e-8) or
                not np.allclose(data.std(axis=0, ddof=1), source_sd, atol=1e-8)):
            raise ValueError(f"Participant cells do not reconcile with sheet summary: {sheet_name}")
        muscles[muscle] = {
            "participants": len(data),
            "frontBarbellConcentricMeanPctReference": rounded(data[:, 0].mean()),
            "frontBarbellEccentricMeanPctReference": rounded(data[:, 1].mean()),
            "frontConcentricMinusEccentric": paired(data[:, 0], data[:, 1]),
            "frontMinusBehindBarbellConcentric": paired(data[:, 0], data[:, 2]),
            "frontBarbellMinusFrontMachineConcentric": paired(data[:, 0], data[:, 4]),
        }
    return {
        "status": "research_seated_press_phase_mean_emg_not_standing_press_or_muscle_force_validation",
        "source": "Coratella et al. 2022, doi:10.3389/fphys.2022.825880; Zenodo doi:10.5281/zenodo.5744051",
        "sourceWorkbookMd5": actual_md5,
        "sourceLicense": "CC BY 4.0 per Zenodo API record 5744051",
        "protocol": "Eight participants, seated front barbell military press with approximately 80-degree back support at 80% measured exercise 1RM; six repetitions, 2 s concentric and 2 s eccentric plus brief isometric phase. Workbook rows 2-9 contain one participant-level normalized EMG mean per exercise and phase; source rows 12-13 reconcile as mean and sample SD. This audit makes within-person paired comparisons without pretending eight means are time series.",
        "unit": "percent of each muscle's maximal voluntary isometric excitation RMS (phase-averaged nRMS); differences are percentage points",
        "limitations": [
            "The source press was seated with a backrest, not the app's standing strict press.",
            "The workbook has only phase averages, no raw synchronized EMG, joint kinematics, barbell motion, hand force, bench contact or individual muscle force.",
            "Electrical activity is not muscle tension, muscle force, or a validated personal excitation prediction.",
            "Eight participants are insufficient to infer a universal muscle ordering or user-specific highlighting.",
        ],
        "muscles": muscles,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()
    report = inspect(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n",
                           encoding="utf-8")
    print(json.dumps(report["muscles"], indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
